"""Extract the 12 beef-cut sheets of the Dataverse xlsx into per-cut CSVs.

Source of record: `e-nose_dataset_12_beef_cuts.xlsx` (Harvard Dataverse
DOI 10.7910/DVN/XNFVTS, CC0 1.0). Each sheet is one beef cut with 2220
minute-by-minute rows: Minute, TVC (log10 CFU/g), Label (1-4, a recoding of
TVC against microbiological thresholds), and 11 MQ sensor readings.

The xlsx `Label` cells are Excel formulas; the cached values are read with
`data_only=True`. Run from this directory:

    python extract.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
XLSX = HERE / "e-nose_dataset_12_beef_cuts.xlsx"
OUT = HERE / "cuts"

HEADER = ["Minute", "TVC", "Label",
          "MQ135", "MQ136", "MQ137", "MQ138", "MQ2", "MQ3",
          "MQ4", "MQ5", "MQ6", "MQ8", "MQ9"]


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    total_rows = 0
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if list(rows[0]) != HEADER:
            raise ValueError(f"{name}: unexpected header {rows[0]}")
        lines = [",".join(HEADER)]
        for r in rows[1:]:
            if any(v is None for v in r):
                raise ValueError(f"{name}: missing value in row {r}")
            lines.append(",".join(str(v) for v in r))
        path = OUT / f"{name}.csv"
        path.write_text("\n".join(lines) + "\n")
        total_rows += len(rows) - 1
        print(f"{name}: {len(rows) - 1} rows -> {path.name}")
    print(f"total rows: {total_rows} (expect 12 x 2220 = 26640)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
